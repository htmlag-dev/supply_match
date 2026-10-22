"""Supplier catalogues. Unknown price/stock stay null; no model invents them."""
import csv
import io
import json
import re
import uuid
from datetime import datetime,timezone
from urllib.parse import urlsplit
from app.services.cost_service import CURRENCIES,number,rounded

class CatalogService:
    def __init__(self,db):
        self.db=db
        with db._conn() as c:
            c.execute('CREATE TABLE IF NOT EXISTS catalog_items(id TEXT PRIMARY KEY,supplier_id INTEGER NOT NULL,payload TEXT NOT NULL)')
            c.execute('CREATE INDEX IF NOT EXISTS catalog_supplier ON catalog_items(supplier_id)')
    @staticmethod
    def decode(row):
        item={**json.loads(row['payload']),'id':row['id'],'supplier_id':row['supplier_id']}
        item.setdefault('checked_at','')
        try:item['stale']=not item['checked_at'] or (datetime.now(timezone.utc)-datetime.fromisoformat(item['checked_at'].replace('Z','+00:00'))).total_seconds()>7*86400
        except (ValueError,TypeError):item['stale']=True
        return item
    def list(self,sid=None):
        with self.db._conn() as c:
            rows=c.execute('SELECT ci.* FROM catalog_items ci JOIN suppliers s ON s.id=ci.supplier_id'+(' WHERE ci.supplier_id=?' if sid is not None else '')+' ORDER BY ci.rowid DESC',(sid,) if sid is not None else ()).fetchall()
            return [self.decode(r) for r in rows]
    def get(self,item_id):
        with self.db._conn() as c:
            row=c.execute('SELECT ci.* FROM catalog_items ci JOIN suppliers s ON s.id=ci.supplier_id WHERE ci.id=?',(item_id,)).fetchone()
            return self.decode(row) if row else None
    def for_ids(self,ids):
        if not ids:return []
        with self.db._conn() as c:
            rows=c.execute('SELECT ci.* FROM catalog_items ci JOIN suppliers s ON s.id=ci.supplier_id WHERE ci.id IN ('+','.join('?' for _ in ids)+')',ids).fetchall()
            return [self.decode(r) for r in rows]
    def groups(self):
        with self.db._conn() as c:
            return [dict(r) for r in c.execute("SELECT DISTINCT ci.supplier_id,json_extract(ci.payload,'$.currency') currency FROM catalog_items ci JOIN suppliers s ON s.id=ci.supplier_id ORDER BY ci.supplier_id,currency")]
    def clean(self,data):
        sid=data.get('supplier_id')
        if isinstance(sid,bool) or not isinstance(sid,int) or not self.db.get_supplier(sid):raise ValueError('Выберите поставщика из базы.')
        item={}
        for key,limit in [('name',300),('category',100),('sku',120),('comparison_key',200),('source_url',2000),('note',1000)]:
            value=data.get(key,'')
            if not isinstance(value,str) or len(value)>limit:raise ValueError('Некорректное поле каталога: '+key)
            item[key]=value.strip()
        if not item['name']:raise ValueError('Укажите точную модель или название товара.')
        item['category']=item['category'] or 'Без категории'
        if item['source_url']:
            parsed=urlsplit(item['source_url'])
            if parsed.scheme not in ('http','https') or not parsed.hostname or parsed.username or parsed.password:raise ValueError('Источник должен быть публичной ссылкой http/https.')
        currency=data.get('currency','RUB')
        if currency not in CURRENCIES:raise ValueError('Неизвестная валюта.')
        item['currency']=currency
        item['unit_price']=None if data.get('unit_price') in (None,'') else rounded(number(data['unit_price'],'Цена'))
        item['availability']=data.get('availability') or 'unknown'
        if item['availability'] not in ('available','out','unknown'):raise ValueError('Неизвестный статус наличия.')
        for key,label in [('stock_qty','Остаток'),('moq','Минимальное количество')]:item[key]=None if data.get(key) in (None,'') else str(number(data[key],label))
        for key in ('lead_min','lead_max'):
            value=data.get(key)
            if value in ('',None):item[key]=None
            else:
                if not re.fullmatch(r'\d{1,4}',str(value)) or int(value)>3650:raise ValueError('Срок поставки — от 0 до 3650 календарных дней.')
                item[key]=int(value)
        if (item['lead_min'] is None)!=(item['lead_max'] is None):raise ValueError('Укажите оба срока поставки или оставьте оба пустыми.')
        if item['lead_min'] is not None and item['lead_min']>item['lead_max']:raise ValueError('Минимальный срок не может превышать максимальный.')
        checked=data.get('checked_at','')
        if not isinstance(checked,str):raise ValueError('Некорректная дата проверки.')
        if checked:
            try:
                stamp=datetime.fromisoformat(checked.replace('Z','+00:00'))
                if stamp.tzinfo is None or stamp>datetime.now(timezone.utc):raise ValueError()
            except ValueError:raise ValueError('Дата проверки должна быть в прошлом с часовым поясом.') from None
        item['checked_at']=checked
        item['updated_at']=datetime.now(timezone.utc).isoformat();item['supplier_id']=sid
        return item
    def save(self,data,item_id=None):
        item=self.clean(data)
        if item_id and 'checked_at' not in data:
            previous=self.get(item_id)
            item['checked_at']=(previous or {}).get('checked_at','') if previous and all(previous.get(k)==item.get(k) for k in ('unit_price','currency','availability','stock_qty','moq','lead_min','lead_max','source_url')) else ''
        with self.db._conn() as c:
            if item_id and not c.execute('SELECT 1 FROM catalog_items WHERE id=? AND supplier_id=?',(item_id,item['supplier_id'])).fetchone():raise ValueError('Позиция не найдена.')
            if not item_id and c.execute('SELECT COUNT(*) FROM catalog_items').fetchone()[0]>=3000:raise ValueError('Лимит каталога — 3000 позиций.')
            item_id=item_id or uuid.uuid4().hex
            old=c.execute('SELECT payload FROM catalog_items WHERE id=?',(item_id,)).fetchone()
            if old:
                old=json.loads(old['payload'])
                if any(old.get(k)!=item.get(k) for k in ('unit_price','currency','availability','stock_qty')):
                    from app.services.procurement_service import ProcurementService
                    # The notification schema is initialized by the API, or lazily here.
                    c.execute('CREATE TABLE IF NOT EXISTS notifications(id TEXT PRIMARY KEY,title TEXT NOT NULL,payload TEXT NOT NULL,created_at INTEGER NOT NULL,is_read INTEGER NOT NULL DEFAULT 0)')
                    ProcurementService.notify(c,'Изменились условия: '+item['name'],{'item_id':item_id,'before':{k:old.get(k) for k in ('unit_price','currency','availability')},'after':{k:item.get(k) for k in ('unit_price','currency','availability')}})
            c.execute('INSERT INTO catalog_items VALUES(?,?,?) ON CONFLICT(id) DO UPDATE SET payload=excluded.payload',(item_id,item['supplier_id'],json.dumps(item,ensure_ascii=False)))
        return {**item,'id':item_id}
    def delete(self,item_id):
        with self.db._conn() as c:c.execute('DELETE FROM catalog_items WHERE id=?',(item_id,))
    def import_csv(self,sid,text):
        return self.import_price_list({'supplier_id':sid,'csv':text},commit=True)

    def import_price_list(self,data,commit=False):
        import base64
        import hashlib
        sid=data.get('supplier_id')
        if data.get('format','csv')=='xlsx':
            import zipfile
            from openpyxl import load_workbook
            encoded=data.get('file','')
            if not isinstance(encoded,str) or len(encoded)>170000:raise ValueError('XLSX должен быть не больше 120 KB.')
            try:
                raw=base64.b64decode(encoded,validate=True)
                with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                    if len(archive.infolist())>200 or sum(i.file_size for i in archive.infolist())>5000000:raise ValueError()
                book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
                try:
                    sheet=book.worksheets[0]
                    if sheet.max_row and sheet.max_row>501:raise ValueError()
                    if sheet.max_column and sheet.max_column>30:raise ValueError()
                    values=list(sheet.iter_rows(min_row=1,max_row=502,max_col=30))
                    if any(cell.data_type=='f' for row in values for cell in row):raise ValueError()
                    rows_values=[[cell.value for cell in row] for row in values]
                finally:book.close()
                while rows_values and all(v is None for v in rows_values[-1]):rows_values.pop()
                headers=[str(v).strip() if v is not None else '' for v in rows_values[0]]
                rows=[{k:('' if v is None else str(v)) for k,v in zip(headers,row) if k} for row in rows_values[1:] if any(v is not None for v in row)]
            except Exception:raise ValueError('Не удалось прочитать XLSX. Нужен первый лист, до 500 строк без формул и до 30 столбцов.') from None
        else:
            value=data.get('csv')
            if not isinstance(value,str) or len(value)>150000:raise ValueError('CSV должен быть не больше 150 KB.')
            try:
                reader=csv.DictReader(io.StringIO(value.lstrip('\ufeff')),delimiter=';' if ';' in value.splitlines()[0] else ',')
                headers=reader.fieldnames or [];rows=list(reader)
            except Exception:raise ValueError('Не удалось прочитать CSV.') from None
        if not rows or len(rows)>500 or 'name' not in headers:raise ValueError('Нужен столбец name и от 1 до 500 строк.')
        nonempty=[h for h in headers if h]
        if len(nonempty)!=len(set(nonempty)):raise ValueError('Повторяющиеся имена столбцов.')
        cleaned=[]
        seen=set()
        for index,row in enumerate(rows,2):
            if None in row:raise ValueError('Лишние поля в строке '+str(index))
            item=self.clean({**row,'supplier_id':sid})
            if item['sku']:
                if item['sku'] in seen:raise ValueError('Повторяющийся артикул в строке '+str(index))
                seen.add(item['sku'])
            cleaned.append(item)
        with self.db._conn() as c:
            existing=[self.decode(r) for r in c.execute('SELECT * FROM catalog_items WHERE supplier_id=?',(sid,))]
            bysku={}
            for item in existing:
                if item['sku']:bysku.setdefault(item['sku'],[]).append(item)
            changes=[]
            for item in cleaned:
                matches=bysku.get(item['sku'],[]) if item['sku'] else []
                if len(matches)>1:raise ValueError('В каталоге несколько позиций с артикулом '+item['sku']+'. Устраните дубликаты.')
                old=matches[0] if matches else None
                # Updating a price list does not certify a human verification.
                if old and 'checked_at' not in headers:item['checked_at']=old.get('checked_at','') if all(old.get(k)==item.get(k) for k in ('unit_price','currency','availability','stock_qty','moq','lead_min','lead_max','source_url')) else ''
                changes.append({'action':'update' if old else 'add','id':old['id'] if old else None,'item':item})
            added=sum(x['action']=='add' for x in changes)
            if c.execute('SELECT COUNT(*) FROM catalog_items').fetchone()[0]+added>3000:raise ValueError('Лимит каталога — 3000 позиций.')
            snapshot=[{k:v for k,v in x.items() if k!='updated_at'} for x in existing]
            signature=hashlib.sha256(json.dumps({'changes':[{**x,'item':{k:v for k,v in x['item'].items() if k!='updated_at'}} for x in changes],'existing':snapshot},sort_keys=True,ensure_ascii=False).encode()).hexdigest()
            if commit:
                if data.get('preview_token') is not None and data['preview_token']!=signature:raise ValueError('Каталог или файл изменился. Повторите предпросмотр.')
                for change in changes:
                    item=change['item'];iid=change['id'] or uuid.uuid4().hex
                    if change['id']:
                        old=next(x for x in existing if x['id']==change['id'])
                        if any(old.get(k)!=item.get(k) for k in ('unit_price','currency','availability','stock_qty')):
                            from app.services.procurement_service import ProcurementService
                            c.execute('CREATE TABLE IF NOT EXISTS notifications(id TEXT PRIMARY KEY,title TEXT NOT NULL,payload TEXT NOT NULL,created_at INTEGER NOT NULL,is_read INTEGER NOT NULL DEFAULT 0)')
                            ProcurementService.notify(c,'Изменились условия: '+item['name'],{'item_id':iid,'before':{k:old.get(k) for k in ('unit_price','currency','availability')},'after':{k:item.get(k) for k in ('unit_price','currency','availability')}})
                    c.execute('INSERT INTO catalog_items VALUES(?,?,?) ON CONFLICT(id) DO UPDATE SET payload=excluded.payload',(iid,sid,json.dumps(item,ensure_ascii=False)))
        return {'added':added,'updated':len(changes)-added,'preview_token':signature,'rows':changes if not commit else []}

    def discover(self,sid,query):
        from ddgs import DDGS
        seller=self.db.get_supplier(sid)
        if not seller:raise ValueError('Поставщик не найден.')
        try:host=urlsplit(seller['site']).hostname
        except ValueError:host=None
        if not host:raise ValueError('У поставщика нет сайта. Добавьте модели вручную или импортируйте каталог.')
        try:results=DDGS(timeout=12).text('site:'+host+' '+str(query or seller['offer'])[:180]+' каталог товар модель',max_results=20)
        except Exception:raise ValueError('Поисковая система временно недоступна. Повторите позже или импортируйте каталог.') from None
        existing={x['source_url'] for x in self.list(sid)};added=0
        for result in results:
            url=str(result.get('href',''))
            try:source=urlsplit(url)
            except ValueError:continue
            if source.scheme not in ('http','https') or not source.hostname or not (source.hostname==host or source.hostname.endswith('.'+host)) or url in existing:continue
            self.save({'supplier_id':sid,'name':str(result.get('title') or host)[:300],'category':'Найденные страницы · проверить модель','source_url':url,'availability':'unknown','note':'Страница из поиска, не подтверждённая товарная карточка. Уточните точную модель, цену и наличие у поставщика.'})
            existing.add(url);added+=1
        return {'added':added,'message':'Найденные страницы добавлены с неизвестными ценой и наличием. Это не полный каталог продавца.'}
